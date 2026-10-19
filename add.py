import random
import threading
import time
import base64
import os
import gradio as gr
from AIGC import GovernmentDocumentSystem
from AIGN import AIGN
from SCRIPT import ScriptAIGN
from zhipuAI import zhipuChatLLM
import re
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import markdown
from bs4 import BeautifulSoup

chatLLM = zhipuChatLLM(model_name="glm-4-plus", api_key='57ecdcf0e2b7407a9800b2e58fc8b214.kXuA501N2Hcc2Q8N')
STREAM_INTERVAL = 0.

# 公文生成常量
STANDARD_FONT = "SimSun"
FONT_MAPPING = {
    "SimSun": ["宋体", "SimSun"],
    "SimHei": ["黑体", "SimHei"],
    "FangSong": ["仿宋", "FangSong_GB2312", "FangSong"]
}

DOC_TYPES = {
    "通知": {
        "structure": ["标题", "发文字号", "主送机关", "正文", "发文机关", "成文日期"],
        "format_rules": {
            "title": {"font": "SimHei", "size": 22, "align": "center"},
            "body": {"font": STANDARD_FONT, "size": 16}
        }
    },
    "请示": {
        "structure": ["标题", "主送机关", "请示缘由", "请示事项", "结语"],
        "format_rules": {
            "title": {"font": "SimHei", "size": 22},
            "emphasis": {"bold": True, "underline": True}
        }
    },
    "函": {
        "structure": ["标题", "主送机关", "发函缘由", "商洽事项", "结语"],
        "format_rules": {
            "title": {"font": "SimHei", "size": 20}
        }
    },
    "报告": {
        "structure": ["标题", "主送机关", "基本情况", "分析情况", "意见建议"],
        "format_rules": {
            "title": {"font": "SimHei", "size": 20}
        }
    }
}

def convert_history(history):
    converted = []
    for msg in history:
        if isinstance(msg, (list, tuple)) and len(msg) == 2:
            converted.append(msg)
        elif isinstance(msg, dict):
            content = msg.get("content", "")
            if msg.get("role") == "user":
                converted.append([None, content])
            elif msg.get("role") == "assistant":
                converted.append([content, None])
    return converted

def make_middle_chat():
    carrier = threading.Event()
    carrier.history = []
    def middle_chat(messages, temperature=None, top_p=None):
        nonlocal carrier
        carrier.history.append([None, ""])
        if len(carrier.history) > 20:
            carrier.history = carrier.history[-16:]
        try:
            output_text = ''
            for resp in chatLLM(messages, temperature=temperature, top_p=top_p, stream=True):
                output_text += resp["content"]
                carrier.history[-1] = [output_text, None]
            return {"content": output_text}
        except Exception as e:
            carrier.history[-1] = [f"Error: {e}", None]
            raise e
    return carrier, middle_chat

def get_available_font(font_name: str) -> str:
    for font in FONT_MAPPING.get(font_name, [font_name]):
        return font
    return "Arial Unicode MS"

def sanitize_text(text: str) -> str:
    text = re.sub(r'[\u3000\xa0\u200b]', ' ', text)
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
    return text.strip()

def format_chinese_date():
    return datetime.now().strftime("%Y年%m月%d日")

def _generate_doc_number(doc_type: str) -> str:
    prefix_map = {"通知": "X政", "请示": "X政请", "函": "X政函", "报告": "X政报"}
    year = datetime.now().year
    seq = f"{int(time.time()) % 1000:03d}"
    return f"{prefix_map.get(doc_type, 'X政')}〔{year}〕{seq}号"

def _set_font(element, font_name: str):
    rPr = element._element.rPr
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        element._element.insert(0, rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)

def generate_document(doc_type, sender, recipient, content, doc_number, doc_date):
    try:
        full_input = {
            "type": doc_type,
            "sender": sender,
            "recipient": recipient,
            "content": content,
            "doc_number": doc_number or _generate_doc_number(doc_type),
            "date": doc_date or format_chinese_date()
        }
        
        messages = [
            {"role": "system", "content": f"请按标准格式生成{doc_type}，直接返回最终公文内容"},
            {"role": "user", "content": content}
        ]
        
        response = chatLLM(messages)
        result = response["content"]
        
        lines = result.split('\n')
        formatted = []
        for i, line in enumerate(lines):
            if i in [0, 1, 3] or i >= len(lines)-2:
                formatted.append(line)
            else:
                if line.strip():
                    formatted.append(f"　　{line}" if not line.startswith(' ') else line)
        
        return "\n".join(formatted)
    except Exception as e:
        return f"生成失败: {str(e)}"

def export_to_word(content_text, doc_type):
    try:
        lines = content_text.split('\n')
        doc_info = {
            "type": doc_type,
            "title": lines[0] if lines else "",
            "doc_number": lines[1] if len(lines)>1 else "",
            "recipient": lines[3] if len(lines)>3 else "",
            "content": '\n'.join(lines[4:-2]) if len(lines)>5 else "",
            "sender": lines[-2] if len(lines)>=2 else "",
            "date": lines[-1] if lines else ""
        }
        
        output_path = os.path.join("output", f"{doc_type}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx")
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        
        document = Document()
        section = document.sections[0]
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.5)
        
        # 标题
        title_para = document.add_paragraph()
        title_run = title_para.add_run(sanitize_text(doc_info["title"]))
        _set_font(title_run, get_available_font("SimHei"))
        title_run.font.size = Pt(22)
        title_run.font.bold = True
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph()
        
        # 文号
        if doc_info.get("doc_number"):
            num_para = document.add_paragraph()
            num_run = num_para.add_run(sanitize_text(doc_info["doc_number"]))
            _set_font(num_run, get_available_font(STANDARD_FONT))
            num_run.font.size = Pt(16)
            num_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            document.add_paragraph()
        
        # 正文内容
        body_font = get_available_font(STANDARD_FONT)
        body_size = Pt(16)
        
        content = sanitize_text(doc_info["content"])
        for paragraph in content.split('\n'):
            if not paragraph.strip():
                document.add_paragraph()
                continue
                
            p = document.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(32)
            p.paragraph_format.line_spacing = 1.5
            
            run = p.add_run(paragraph)
            _set_font(run, body_font)
            run.font.size = body_size

        # 发文信息
        if all(key in doc_info for key in ["sender", "date"]):
            document.add_paragraph()
            
            sender_para = document.add_paragraph()
            sender_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            sender_run = sender_para.add_run(sanitize_text(doc_info["sender"]))
            _set_font(sender_run, body_font)
            sender_run.font.size = body_size
            
            date_para = document.add_paragraph()
            date_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            date_run = date_para.add_run(sanitize_text(doc_info["date"]))
            _set_font(date_run, body_font)
            date_run.font.size = body_size

        document.save(output_path)
        return output_path
    except Exception as e:
        raise gr.Error(f"Word导出失败: {str(e)}")

def gen_ouline_button_clicked(aign, user_idea, history, chapter_count, word_count_per_chapter):
    history = convert_history(history)
    aign.user_idea = user_idea
    aign.chapter_count = chapter_count
    aign.word_count_per_chapter = word_count_per_chapter

    carrier, middle_chat = make_middle_chat()
    carrier.history = history
    aign.novel_outline_writer.chatLLM = middle_chat

    gen_ouline_thread = threading.Thread(target=aign.genNovelOutline)
    gen_ouline_thread.start()

    while gen_ouline_thread.is_alive():
        yield [
            aign,
            carrier.history,
            aign.novel_outline,
            gr.Button(visible=True),
            gr.Button(visible=False),
            gr.Button(visible=False),
            gr.Button(visible=False)
        ]
        time.sleep(STREAM_INTERVAL)
        
    yield [
        aign,
        carrier.history,
        aign.novel_outline,
        gr.Button(visible=True),
        gr.Button(visible=True),
        gr.Button(visible=False),
        gr.Button(visible=False)
    ]

def gen_beginning_button_clicked(aign, history, novel_outline, user_requriments, embellishment_idea):
    history = convert_history(history)
    aign.novel_outline = novel_outline
    aign.user_requriments = user_requriments
    aign.embellishment_idea = embellishment_idea

    carrier, middle_chat = make_middle_chat()
    carrier.history = history
    aign.novel_beginning_writer.chatLLM = middle_chat
    aign.novel_embellisher.chatLLM = middle_chat

    gen_beginning_thread = threading.Thread(target=aign.genBeginning)
    gen_beginning_thread.start()

    while gen_beginning_thread.is_alive():
        yield [
            aign,
            carrier.history,
            aign.writing_plan,
            aign.temp_setting,
            aign.novel_content,
            gr.Button(visible=True),
            gr.Button(visible=True),
            gr.Button(visible=False)
        ]
        time.sleep(STREAM_INTERVAL)
        
    yield [
        aign,
        carrier.history,
        aign.writing_plan,
        aign.temp_setting,
        aign.novel_content,
        gr.Button(visible=True),
        gr.Button(visible=True),
        gr.Button(visible=True)
    ]

def gen_next_paragraph_button_clicked(
    aign,
    history,
    user_idea,
    novel_outline,
    writing_memory,
    temp_setting,
    writing_plan,
    user_requriments,
    embellishment_idea,
    auto_gen_next
):
    history = convert_history(history)
    current_state = {
        "paragraph_list": aign.paragraph_list.copy(),
        "writing_plan": aign.writing_plan,
        "temp_setting": aign.temp_setting,
        "writing_memory": aign.writing_memory,
        "novel_content": aign.novel_content
    }
    aign.history_stack.append(current_state)
    if len(aign.history_stack) > aign.max_history:
        aign.history_stack.pop(0)
        
    aign.user_idea = user_idea
    aign.novel_outline = novel_outline
    aign.writing_memory = writing_memory
    aign.temp_setting = temp_setting
    aign.writing_plan = writing_plan
    aign.user_requriments = user_requriments
    aign.embellishment_idea = embellishment_idea

    carrier, middle_chat = make_middle_chat()
    carrier.history = history
    aign.novel_writer.chatLLM = middle_chat
    aign.novel_embellisher.chatLLM = middle_chat
    aign.memory_maker.chatLLM = middle_chat

    # 获取当前章节信息
    current_chapter = len(aign.paragraph_list) + 1
    total_chapters = aign.chapter_count
    
    gen_next_paragraph_thread = threading.Thread(target=aign.genNextParagraph)
    gen_next_paragraph_thread.start()

    while gen_next_paragraph_thread.is_alive():
        yield [
            aign,
            carrier.history,
            aign.writing_plan,
            aign.temp_setting,
            aign.writing_memory,
            aign.novel_content,
            gr.Button(visible=True),
            gr.Button(visible=True),
            gr.Button(visible=True)
        ]
        time.sleep(STREAM_INTERVAL)
        
    # 添加章节标记
    if aign.novel_content and f"第{current_chapter}章" not in aign.novel_content:
        chapter_header = f"\n\n# 第{current_chapter}章/{total_chapters}章\n\n"
        aign.novel_content = aign.novel_content + chapter_header + aign.paragraph_list[-1]
    
    yield [
        aign,
        carrier.history,
        aign.writing_plan,
        aign.temp_setting,
        aign.writing_memory,
        aign.novel_content,
        gr.Button(visible=True),
        gr.Button(visible=True),
        gr.Button(visible=True)
    ]

def gen_ending_button_clicked(aign, history, user_requriments, novel_content):
    history = convert_history(history)
    aign.user_requriments = user_requriments
    
    carrier, middle_chat = make_middle_chat()
    carrier.history = history
    aign.novel_ending_writer.chatLLM = middle_chat
    aign.novel_embellisher.chatLLM = middle_chat

    gen_ending_thread = threading.Thread(target=aign.genEnding)
    gen_ending_thread.start()

    while gen_ending_thread.is_alive():
        yield [
            aign,
            carrier.history,
            aign.novel_content,
            gr.Button(visible=True)
        ]
        time.sleep(STREAM_INTERVAL)
        
    yield [
        aign,
        carrier.history,
        aign.novel_content,
        gr.Button(visible=True)
    ]

def undo_last_action(aign, history):
    if not aign.history_stack:
        raise gr.Error("没有可撤销的操作")
    
    prev_state = aign.history_stack.pop()
    aign.paragraph_list = prev_state["paragraph_list"]
    aign.writing_plan = prev_state["writing_plan"]
    aign.temp_setting = prev_state["temp_setting"]
    aign.writing_memory = prev_state["writing_memory"]
    aign.novel_content = prev_state["novel_content"]
    
    return [
        aign,
        history,
        aign.writing_plan,
        aign.temp_setting,
        aign.writing_memory,
        aign.novel_content
    ]

def export_novel_to_markdown(novel_content, title="小说"):
    try:
        output_dir = os.path.join("output", "novels")
        os.makedirs(output_dir, exist_ok=True)
        
        filename = f"{title}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.md"
        output_path = os.path.join(output_dir, filename)
        
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(f"# {title}\n\n")
            f.write(novel_content)
            
        return output_path
    except Exception as e:
        raise gr.Error(f"Markdown导出失败: {str(e)}")

def export_novel_to_word(novel_content, title="小说"):
    try:
        output_dir = os.path.join("output", "novels")
        os.makedirs(output_dir, exist_ok=True)
        
        filename = f"{title}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        output_path = os.path.join(output_dir, filename)
        
        document = Document()
        section = document.sections[0]
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)
        
        # 标题
        title_para = document.add_paragraph()
        title_run = title_para.add_run(title)
        _set_font(title_run, get_available_font("SimHei"))
        title_run.font.size = Pt(22)
        title_run.font.bold = True
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph()
        
        # 转换Markdown为Word格式
        html = markdown.markdown(novel_content)
        soup = BeautifulSoup(html, 'html.parser')
        
        for element in soup.children:
            if element.name == "h1":
                p = document.add_paragraph()
                p.paragraph_format.space_after = Pt(12)
                run = p.add_run(element.get_text())
                _set_font(run, get_available_font("SimHei"))
                run.font.size = Pt(18)
                run.font.bold = True
            elif element.name == "h2":
                p = document.add_paragraph()
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run(element.get_text())
                _set_font(run, get_available_font("SimHei"))
                run.font.size = Pt(16)
                run.font.bold = True
            elif element.name == "p":
                p = document.add_paragraph()
                p.paragraph_format.first_line_indent = Pt(32)
                p.paragraph_format.line_spacing = 1.5
                run = p.add_run(element.get_text())
                _set_font(run, get_available_font(STANDARD_FONT))
                run.font.size = Pt(14)
            elif element.name == "blockquote":
                p = document.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.5)
                p.paragraph_format.first_line_indent = Inches(-0.5)
                run = p.add_run(element.get_text())
                _set_font(run, get_available_font("FangSong"))
                run.font.size = Pt(13)
                run.font.italic = True
        
        document.save(output_path)
        return output_path
    except Exception as e:
        raise gr.Error(f"Word导出失败: {str(e)}")

def generate_cover_image(prompt, style="水墨风格"):
    try:
        from diffusers import StableDiffusionPipeline
        import torch
        
        model_id = "runwayml/stable-diffusion-v1-5"
        pipe = StableDiffusionPipeline.from_pretrained(model_id, torch_dtype=torch.float16)
        pipe = pipe.to("cuda")
        
        full_prompt = f"{prompt}, {style}, 高清, 8k, 精致的细节"
        image = pipe(full_prompt).images[0]
        
        output_dir = os.path.join("output", "covers")
        os.makedirs(output_dir, exist_ok=True)
        
        filename = f"cover_{int(time.time())}.png"
        output_path = os.path.join(output_dir, filename)
        image.save(output_path)
        
        return output_path
    except Exception as e:
        raise gr.Error(f"封面生成失败: {str(e)}")

def generate_music(prompt, duration=30, style="中国风"):
    try:
        from audiocraft.models import MusicGen
        from audiocraft.data.audio import audio_write
        
        model = MusicGen.get_pretrained("melody")
        model.set_generation_params(duration=duration)
        
        descriptions = [f"{prompt}, {style}音乐"]
        wav = model.generate(descriptions)
        
        output_dir = os.path.join("output", "music")
        os.makedirs(output_dir, exist_ok=True)
        
        filename = f"music_{int(time.time())}.wav"
        output_path = os.path.join(output_dir, filename)
        audio_write(output_path, wav[0].cpu(), model.sample_rate, strategy="loudness")
        
        return output_path
    except Exception as e:
        raise gr.Error(f"音乐生成失败: {str(e)}")

def create_gradio_interface():
    with gr.Blocks(title="大国AI创作助手", theme=gr.themes.Soft()) as demo:
        gr.Markdown("""
        <div style="text-align: center;">
            <h1>大国AI创作助手</h1>
            <p>全能型AI创作平台，支持公文写作、小说创作、AI绘画和音乐生成</p>
        </div>
        """)
        
        with gr.Tabs():
            # 公文写作标签页
            with gr.Tab("公文写作"):
                with gr.Row():
                    with gr.Column(scale=1):
                        doc_type = gr.Dropdown(
                            label="公文类型",
                            choices=["通知", "请示", "函", "报告"],
                            value="通知"
                        )
                        sender = gr.Textbox(label="发文机关", placeholder="请输入发文机关名称")
                        recipient = gr.Textbox(label="主送机关", placeholder="请输入主送机关名称")
                        doc_number = gr.Textbox(label="发文字号(可选)", placeholder="留空自动生成")
                        doc_date = gr.Textbox(label="成文日期(可选)", placeholder="留空使用当前日期")
                        
                    with gr.Column(scale=2):
                        content = gr.Textbox(
                            label="正文内容",
                            placeholder="请输入公文主要内容...",
                            lines=10,
                            max_lines=20
                        )
                        with gr.Row():
                            generate_btn = gr.Button("生成公文", variant="primary")
                            export_word_btn = gr.Button("导出Word", variant="secondary")
                        
                        output = gr.Textbox(label="生成的公文", interactive=False, lines=15)
                
                generate_btn.click(
                    generate_document,
                    inputs=[doc_type, sender, recipient, content, doc_number, doc_date],
                    outputs=output
                )
                
                export_word_btn.click(
                    export_to_word,
                    inputs=[output, doc_type],
                    outputs=gr.File(label="下载公文")
                )
            
            # 小说创作标签页
            with gr.Tab("小说创作"):
                with gr.Row():
                    with gr.Column(scale=1):
                        user_idea = gr.Textbox(
                            label="创作灵感",
                            placeholder="请输入你的小说创意...",
                            lines=3
                        )
                        chapter_count = gr.Slider(
                            label="章节数量",
                            minimum=1,
                            maximum=50,
                            step=1,
                            value=10
                        )
                        word_count_per_chapter = gr.Slider(
                            label="每章字数",
                            minimum=500,
                            maximum=5000,
                            step=100,
                            value=2000
                        )
                        user_requriments = gr.Textbox(
                            label="特殊要求",
                            placeholder="如人物设定、风格要求等...",
                            lines=2
                        )
                        embellishment_idea = gr.Textbox(
                            label="润色要求",
                            placeholder="对文笔风格的具体要求...",
                            lines=2
                        )
                        
                    with gr.Column(scale=2):
                        chatbot = gr.Chatbot(
                            label="创作过程",
                            height=400,
                            bubble_full_width=False
                        )
                        with gr.Row():
                            gen_outline_btn = gr.Button("生成大纲", variant="primary")
                            gen_beginning_btn = gr.Button("生成开头", variant="primary", visible=False)
                            gen_next_btn = gr.Button("续写下一段", variant="primary", visible=False)
                            gen_ending_btn = gr.Button("生成结尾", variant="primary", visible=False)
                            undo_btn = gr.Button("撤销", variant="secondary")
                        
                        novel_outline = gr.Textbox(
                            label="小说大纲",
                            interactive=False,
                            lines=8,
                            visible=False
                        )
                        writing_plan = gr.Textbox(
                            label="写作计划",
                            interactive=False,
                            lines=4,
                            visible=False
                        )
                        temp_setting = gr.Textbox(
                            label="临时设定",
                            interactive=False,
                            lines=4,
                            visible=False
                        )
                        writing_memory = gr.Textbox(
                            label="创作记忆",
                            interactive=False,
                            lines=4,
                            visible=False
                        )
                        novel_content = gr.Textbox(
                            label="小说内容",
                            interactive=True,
                            lines=15
                        )
                        with gr.Row():
                            export_md_btn = gr.Button("导出Markdown", variant="secondary")
                            export_word_btn = gr.Button("导出Word", variant="secondary")
                            gen_cover_btn = gr.Button("生成封面", variant="secondary")
                            gen_music_btn = gr.Button("生成配乐", variant="secondary")
                
                aign = gr.State(AIGN(chatLLM))
                
                gen_outline_btn.click(
                    gen_ouline_button_clicked,
                    inputs=[aign, user_idea, chatbot, chapter_count, word_count_per_chapter],
                    outputs=[aign, chatbot, novel_outline, gen_beginning_btn, gen_next_btn, gen_ending_btn, undo_btn]
                )
                
                gen_beginning_btn.click(
                    gen_beginning_button_clicked,
                    inputs=[aign, chatbot, novel_outline, user_requriments, embellishment_idea],
                    outputs=[aign, chatbot, writing_plan, temp_setting, novel_content, gen_outline_btn, gen_next_btn, gen_ending_btn]
                )
                
                gen_next_btn.click(
                    gen_next_paragraph_button_clicked,
                    inputs=[
                        aign, chatbot, user_idea, novel_outline, writing_memory,
                        temp_setting, writing_plan, user_requriments, embellishment_idea,
                        gr.Checkbox(False, visible=False)
                    ],
                    outputs=[aign, chatbot, writing_plan, temp_setting, writing_memory, novel_content, gen_outline_btn, gen_beginning_btn, gen_ending_btn]
                )
                
                gen_ending_btn.click(
                    gen_ending_button_clicked,
                    inputs=[aign, chatbot, user_requriments, novel_content],
                    outputs=[aign, chatbot, novel_content, undo_btn]
                )
                
                undo_btn.click(
                    undo_last_action,
                    inputs=[aign, chatbot],
                    outputs=[aign, chatbot, writing_plan, temp_setting, writing_memory, novel_content]
                )
                
                export_md_btn.click(
                    export_novel_to_markdown,
                    inputs=[novel_content, user_idea],
                    outputs=gr.File(label="下载小说")
                )
                
                export_word_btn.click(
                    export_novel_to_word,
                    inputs=[novel_content, user_idea],
                    outputs=gr.File(label="下载小说")
                )
                
                gen_cover_btn.click(
                    generate_cover_image,
                    inputs=[user_idea],
                    outputs=gr.Image(label="小说封面", type="filepath")
                )
                
                gen_music_btn.click(
                    generate_music,
                    inputs=[user_idea],
                    outputs=gr.Audio(label="背景音乐", type="filepath")
                )
            
            # AI绘画标签页
            with gr.Tab("AI绘画"):
                with gr.Row():
                    with gr.Column(scale=1):
                        paint_prompt = gr.Textbox(
                            label="绘画描述",
                            placeholder="详细描述你想要绘制的画面...",
                            lines=4
                        )
                        paint_style = gr.Dropdown(
                            label="绘画风格",
                            choices=["水墨风格", "油画风格", "卡通风格", "写实风格", "赛博朋克"],
                            value="水墨风格"
                        )
                        paint_btn = gr.Button("开始绘画", variant="primary")
                        
                    with gr.Column(scale=2):
                        paint_output = gr.Image(label="生成作品", type="filepath")
                        
                paint_btn.click(
                    generate_cover_image,
                    inputs=[paint_prompt, paint_style],
                    outputs=paint_output
                )
            
            # AI音乐标签页
            with gr.Tab("AI音乐"):
                with gr.Row():
                    with gr.Column(scale=1):
                        music_prompt = gr.Textbox(
                            label="音乐描述",
                            placeholder="描述你想要的音乐风格和情感...",
                            lines=4
                        )
                        music_style = gr.Dropdown(
                            label="音乐类型",
                            choices=["中国风", "流行乐", "电子乐", "古典乐", "摇滚乐"],
                            value="中国风"
                        )
                        music_duration = gr.Slider(
                            label="时长(秒)",
                            minimum=10,
                            maximum=120,
                            step=5,
                            value=30
                        )
                        music_btn = gr.Button("生成音乐", variant="primary")
                        
                    with gr.Column(scale=2):
                        music_output = gr.Audio(label="生成音乐", type="filepath")
                        
                music_btn.click(
                    generate_music,
                    inputs=[music_prompt, music_duration, music_style],
                    outputs=music_output
                )
        
        # 底部信息
        gr.Markdown("""
        <div style="text-align: center; margin-top: 20px;">
            <p>大国AI创作助手 - 让创作更简单</p>
            <p>关注微信公众号获取更多功能</p>
        </div>
        """)
    
    return demo

if __name__ == "__main__":
    demo = create_gradio_interface()
    demo.launch(server_name="127.0.0.1", server_port=7860)
