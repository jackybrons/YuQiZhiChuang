import random
import threading
import time
import base64
import os
import gradio as gr
from AIGC import GovernmentDocumentSystem
from AIGN import AIGN
from zhipuAI import zhipuChatLLM
import re
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import markdown
from bs4 import BeautifulSoup

chatLLM = zhipuChatLLM(model_name="glm-4-plus", api_key='57ecdcf0e2b7407a9800b2e58fc8b214.kXuA501N2Hcc2Q8N')
STREAM_INTERVAL = 0.

# 公文生成常量
STANDARD_FONT = "SimSun"
FONT_MAPPING = {
    "SimSun": ["宋体", "SimSun"],
    "SimHei": ["黑体", "SimHei"],
    "FangSong": ["仿宋", "FangSong_GB2312", "FangSong"]
}

DOC_TYPES = {
    "通知": {
        "structure": ["标题", "发文字号", "主送机关", "正文", "发文机关", "成文日期"],
        "format_rules": {
            "title": {"font": "SimHei", "size": 22, "align": "center"},
            "body": {"font": STANDARD_FONT, "size": 16}
        }
    },
    "请示": {
        "structure": ["标题", "主送机关", "请示缘由", "请示事项", "结语"],
        "format_rules": {
            "title": {"font": "SimHei", "size": 22},
            "emphasis": {"bold": True, "underline": True}
        }
    },
    "函": {
        "structure": ["标题", "主送机关", "发函缘由", "商洽事项", "结语"],
        "format_rules": {
            "title": {"font": "SimHei", "size": 20}
        }
    },
    "报告": {
        "structure": ["标题", "主送机关", "基本情况", "分析情况", "意见建议"],
        "format_rules": {
            "title": {"font": "SimHei", "size": 20}
        }
    }
}

def convert_history(history):
    converted = []
    for msg in history:
        if isinstance(msg, (list, tuple)):
            role = "user" if msg[0] is None else "assistant"
            content = msg[1] if msg[0] is None else msg[0]
            converted.append({"role": role, "content": content})
        elif isinstance(msg, dict):
            converted.append(msg)
    return converted

def make_middle_chat():
    carrier = threading.Event()
    carrier.history = []
    def middle_chat(messages, temperature=None, top_p=None):
        nonlocal carrier
        carrier.history.append({"role": "assistant", "content": ""})
        if len(carrier.history) > 20:
            carrier.history = carrier.history[-16:]
        try:
            output_text = ''
            for resp in chatLLM(messages, temperature=temperature, top_p=top_p, stream=True):
                output_text += resp["content"]
                carrier.history[-1] = {
                    "role": "assistant",
                    "content": output_text
                }
            return {"content": output_text}
        except Exception as e:
            carrier.history[-1] = {"role": "assistant", "content": f"Error: {e}"}
            raise e
    return carrier, middle_chat

def get_available_font(font_name: str) -> str:
    for font in FONT_MAPPING.get(font_name, [font_name]):
        return font
    return "Arial Unicode MS"

def sanitize_text(text: str) -> str:
    text = re.sub(r'[\u3000\xa0\u200b]', ' ', text)
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
    return text.strip()

def format_chinese_date():
    return datetime.now().strftime("%Y年%m月%d日")

def _generate_doc_number(doc_type: str) -> str:
    prefix_map = {"通知": "X政", "请示": "X政请", "函": "X政函", "报告": "X政报"}
    year = datetime.now().year
    seq = f"{int(time.time()) % 1000:03d}"
    return f"{prefix_map.get(doc_type, 'X政')}〔{year}〕{seq}号"

def _set_font(element, font_name: str):
    rPr = element._element.rPr
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        element._element.insert(0, rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)

def generate_document(doc_type, sender, recipient, content, doc_number, doc_date):
    try:
        full_input = {
            "type": doc_type,
            "sender": sender,
            "recipient": recipient,
            "content": content,
            "doc_number": doc_number or _generate_doc_number(doc_type),
            "date": doc_date or format_chinese_date()
        }
        
        messages = [
            {"role": "system", "content": f"请按标准格式生成{doc_type}，直接返回最终公文内容"},
            {"role": "user", "content": content}
        ]
        
        response = chatLLM(messages)
        result = response["content"]  # 修改此处，直接获取响应内容
        
        lines = result.split('\n')
        formatted = []
        for i, line in enumerate(lines):
            if i in [0, 1, 3] or i >= len(lines)-2:
                formatted.append(line)
            else:
                if line.strip():
                    formatted.append(f"　　{line}" if not line.startswith(' ') else line)
        
        return "\n".join(formatted)
    except Exception as e:
        return f"生成失败: {str(e)}"

def export_to_word(content_text, doc_type):
    try:
        lines = content_text.split('\n')
        doc_info = {
            "type": doc_type,
            "title": lines[0] if lines else "",
            "doc_number": lines[1] if len(lines)>1 else "",
            "recipient": lines[3] if len(lines)>3 else "",
            "content": '\n'.join(lines[4:-2]) if len(lines)>5 else "",
            "sender": lines[-2] if len(lines)>=2 else "",
            "date": lines[-1] if lines else ""
        }
        
        output_path = os.path.join("output", f"{doc_type}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx")
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        
        document = Document()
        section = document.sections[0]
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.5)
        
        # 标题
        title_para = document.add_paragraph()
        title_run = title_para.add_run(sanitize_text(doc_info["title"]))
        _set_font(title_run, get_available_font("SimHei"))
        title_run.font.size = Pt(22)
        title_run.font.bold = True
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph()
        
        # 文号
        if doc_info.get("doc_number"):
            num_para = document.add_paragraph()
            num_run = num_para.add_run(sanitize_text(doc_info["doc_number"]))
            _set_font(num_run, get_available_font(STANDARD_FONT))
            num_run.font.size = Pt(16)
            num_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            document.add_paragraph()
        
        # 正文内容
        body_font = get_available_font(STANDARD_FONT)
        body_size = Pt(16)
        
        content = sanitize_text(doc_info["content"])
        for paragraph in content.split('\n'):
            if not paragraph.strip():
                document.add_paragraph()
                continue
                
            p = document.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(32)
            p.paragraph_format.line_spacing = 1.5
            
            run = p.add_run(paragraph)
            _set_font(run, body_font)
            run.font.size = body_size

        # 发文信息
        if all(key in doc_info for key in ["sender", "date"]):
            document.add_paragraph()
            
            sender_para = document.add_paragraph()
            sender_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            sender_run = sender_para.add_run(sanitize_text(doc_info["sender"]))
            _set_font(sender_run, body_font)
            sender_run.font.size = body_size
            
            date_para = document.add_paragraph()
            date_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            date_run = date_para.add_run(sanitize_text(doc_info["date"]))
            _set_font(date_run, body_font)
            date_run.font.size = body_size

        document.save(output_path)
        return output_path
    except Exception as e:
        raise gr.Error(f"Word导出失败: {str(e)}")

def gen_ouline_button_clicked(aign, user_idea, history):
    history = convert_history(history)
    aign.user_idea = user_idea

    carrier, middle_chat = make_middle_chat()
    carrier.history = history
    aign.novel_outline_writer.chatLLM = middle_chat

    gen_ouline_thread = threading.Thread(target=aign.genNovelOutline)
    gen_ouline_thread.start()

    while gen_ouline_thread.is_alive():
        yield [
            aign,
            carrier.history,
            aign.novel_outline,
            gr.Button(visible=False),
        ]
        time.sleep(STREAM_INTERVAL)
    yield [
        aign,
        carrier.history,
        aign.novel_outline,
        gr.Button(visible=False),
    ]

def gen_beginning_button_clicked(aign, history, novel_outline, user_requriments, embellishment_idea):
    history = convert_history(history)
    aign.novel_outline = novel_outline
    aign.user_requriments = user_requriments
    aign.embellishment_idea = embellishment_idea

    carrier, middle_chat = make_middle_chat()
    carrier.history = history
    aign.novel_beginning_writer.chatLLM = middle_chat
    aign.novel_embellisher.chatLLM = middle_chat

    gen_beginning_thread = threading.Thread(target=aign.genBeginning)
    gen_beginning_thread.start()

    while gen_beginning_thread.is_alive():
        yield [
            aign,
            carrier.history,
            aign.writing_plan,
            aign.temp_setting,
            aign.novel_content,
            gr.Button(visible=False),
        ]
        time.sleep(STREAM_INTERVAL)
    yield [
        aign,
        carrier.history,
        aign.writing_plan,
        aign.temp_setting,
        aign.novel_content,
        gr.Button(visible=False),
    ]

def gen_next_paragraph_button_clicked(
    aign,
    history,
    user_idea,
    novel_outline,
    writing_memory,
    temp_setting,
    writing_plan,
    user_requriments,
    embellishment_idea,
    auto_gen_next
):
    history = convert_history(history)
    current_state = {
        "paragraph_list": aign.paragraph_list.copy(),
        "writing_plan": aign.writing_plan,
        "temp_setting": aign.temp_setting,
        "writing_memory": aign.writing_memory,
        "novel_content": aign.novel_content
    }
    aign.history_stack.append(current_state)
    if len(aign.history_stack) > aign.max_history:
        aign.history_stack.pop(0)
        
    aign.user_idea = user_idea
    aign.novel_outline = novel_outline
    aign.writing_memory = writing_memory
    aign.temp_setting = temp_setting
    aign.writing_plan = writing_plan
    aign.user_requriments = user_requriments
    aign.embellishment_idea = embellishment_idea

    carrier, middle_chat = make_middle_chat()
    carrier.history = history
    aign.novel_writer.chatLLM = middle_chat
    aign.novel_embellisher.chatLLM = middle_chat
    aign.memory_maker.chatLLM = middle_chat

    gen_next_paragraph_thread = threading.Thread(target=aign.genNextParagraph)
    gen_next_paragraph_thread.start()

    while gen_next_paragraph_thread.is_alive():
        yield [
            aign,
            carrier.history,
            aign.writing_plan,
            aign.temp_setting,
            aign.writing_memory,
            aign.novel_content,
            gr.Button(visible=False),
        ]
        time.sleep(STREAM_INTERVAL)
        
    yield [
        aign,
        carrier.history,
        aign.writing_plan,
        aign.temp_setting,
        aign.writing_memory,
        aign.novel_content,
        gr.Button(visible=False),
    ]

def gen_ending_button_clicked(aign, history, user_requriments, novel_content):
    history = convert_history(history)
    aign.user_requriments = user_requriments
    
    carrier, middle_chat = make_middle_chat()
    carrier.history = history
    aign.novel_ending_writer.chatLLM = middle_chat
    aign.novel_embellisher.chatLLM = middle_chat

    current_state = {
        "paragraph_list": aign.paragraph_list.copy(),
        "writing_plan": aign.writing_plan,
        "temp_setting": aign.temp_setting,
        "writing_memory": aign.writing_memory,
        "novel_content": aign.novel_content
    }
    aign.history_stack.append(current_state)
    if len(aign.history_stack) > aign.max_history:
        aign.history_stack.pop(0)

    gen_ending_thread = threading.Thread(target=aign.genEnding)
    gen_ending_thread.start()

    while gen_ending_thread.is_alive():
        yield [
            aign,
            carrier.history,
            aign.writing_plan,
            aign.temp_setting,
            aign.writing_memory,
            aign.novel_content,
            gr.Button(visible=False),
        ]
        time.sleep(STREAM_INTERVAL)
        
    yield [
        aign,
        carrier.history,
        aign.writing_plan,
        aign.temp_setting,
        aign.writing_memory,
        aign.novel_content,
        gr.Button(visible=False),
    ]

def download_novel():
    if os.path.exists("novel_record.md"):
        with open("novel_record.md", "r+", encoding="utf-8") as f:
            content = f.read()
            if "[END]" not in content:
                f.write("\n\n# [END]\n")
        return "novel_record.md"
    raise gr.Error("尚未生成任何小说内容")

def download_document(doc_type, content):
    try:
        if not content.strip():
            raise gr.Error("没有可下载的内容")
        
        output_path = export_to_word(content, doc_type)
        if os.path.exists(output_path):
            return output_path
        raise gr.Error("文件生成失败")
    except Exception as e:
        raise gr.Error(f"下载失败: {str(e)}")

def toggle_interface(mode):
    is_fiction = mode == "fiction"
    return {
        fiction_interface: gr.Row(visible=is_fiction),
        document_interface: gr.Row(visible=not is_fiction),
        switch_btn: gr.Button(
            value="切换到公文模式" if is_fiction else "切换到小说模式",
            variant="primary" if not is_fiction else "secondary"
        )
    }

def process_document_request(doc_type, doc_requirements, history):
    history = convert_history(history)
    
    try:
        # 初始化公文生成所需参数
        sender = "XX市城市绿化管理局"
        recipient = "各区县城市绿化管理部门、各相关单位"
        doc_number = _generate_doc_number(doc_type)
        doc_date = format_chinese_date()
        
        # 生成公文内容
        generated_content = generate_document(
            doc_type=doc_type,
            sender=sender,
            recipient=recipient,
            content=doc_requirements,
            doc_number=doc_number,
            doc_date=doc_date
        )
        
        # 格式化输出
        formatted_output = f"📝 根据您的要求生成的{doc_type}：\n\n{generated_content}"
        
        # 更新历史记录
        new_history = history + [
            {"role": "user", "content": doc_requirements},
            {"role": "assistant", "content": formatted_output}
        ]
        
        return [new_history, generated_content, gr.Button(visible=True)]
    
    except Exception as e:
        error_msg = f"❌ {doc_type}生成失败：{str(e)}"
        return [
            history + [{"role": "assistant", "content": error_msg}],
            "",
            gr.Button(visible=False)
        ]

css = """
/* 主容器透明化 */
gradio-app {
    background-color: rgba(0,0,0,0) !important;
    background-image: none !important;
}

/* 全局文字颜色 */
* {
    color: black !important;
    background: none transparent !important;
}

body {
    background: #e6f7ff !important;
}

/* 新增大尺寸输入输出框样式 */
.large-input .gr-input,
.large-output .gr-textbox {
    min-height: 600px !important;
    font-size: 14px !important;
    line-height: 1.6 !important;
}

/* 标题行样式 */
.title-row {
    margin-bottom: 20px !important;
}

.mode-button {
    margin-left: 20px !important;
    background: #2196F3 !important;
    color: white !important;
    padding: 6px 12px !important;
    font-size: 14px !important;
    border-radius: 4px !important;
    border: none !important;
    cursor: pointer !important;
}

.document-chat .message-bot {
    font-family: 'SimSun', serif !important;
    white-space: pre-wrap;
    background: #f0faff !important;
    border-left: 4px solid #1890ff !important;
}

.document-preview {
    font-family: 'FangSong', serif !important;
    font-size: 16px !important;
    line-height: 2 !important;
    padding: 20px !important;
    background: #f9f9f9 !important;
    border: 1px solid #ddd !important;
}

.hidden {
    display: none !important;
}
"""

with gr.Blocks(css=css) as demo:
    aign = gr.State(AIGN(chatLLM))
    current_mode = gr.State("fiction")
    
    with gr.Row(equal_height=True, elem_classes="title-row"):
        gr.Markdown("## 语启智创--文本创作助手")
        switch_btn = gr.Button("切换到公文模式", elem_classes="mode-button")
    
    # 小说创作界面
    with gr.Row(visible=True) as fiction_interface:
        with gr.Column(scale=1, elem_id="row1"):
            with gr.Tab("开始"):
                gr.Markdown("生成大纲->生成开头->生成下一段")
                user_idea_text = gr.Textbox(
                    "请在这里输入您的想法",
                    label="想法",
                    lines=7,
                    interactive=True,
                )
                user_requriments_text = gr.Textbox(
                    "",
                    label="写作要求",
                    lines=7,
                    interactive=True,
                )
                embellishment_idea_text = gr.Textbox(
                    "",
                    label="润色要求",
                    lines=7,
                    interactive=True,
                )
                gen_ouline_button = gr.Button("生成大纲")
            with gr.Tab("大纲"):
                novel_outline_text = gr.Textbox(
                    label="大纲", lines=30, interactive=True
                )
                gen_beginning_button = gr.Button("生成开头")
            with gr.Tab("状态"):
                writing_memory_text = gr.Textbox(
                    label="记忆",
                    lines=6,
                    interactive=True,
                    max_lines=8,
                )
                writing_plan_text = gr.Textbox(label="计划", lines=6, interactive=True)
                temp_setting_text = gr.Textbox(
                    label="临时设定", lines=5, interactive=True
                )
                with gr.Row():
                    gen_next_paragraph_button = gr.Button("生成下一段")
                    gen_ending_button = gr.Button("生成结局", variant="stop")
                with gr.Row():
                    download_button = gr.Button("下载小说")

        with gr.Column(scale=3, elem_id="row2"):
            chatBox = gr.Chatbot(height="80vh", label="", type="messages")
        with gr.Column(scale=0, elem_id="row3"):
            novel_content_text = gr.Textbox(
                label="小说正文", lines=32, interactive=True, show_copy_button=True
            )
    
    # 公文创作界面
    with gr.Row(visible=False) as document_interface:
        with gr.Column(scale=1, elem_classes="large-input"):
            doc_type = gr.Dropdown(
                choices=["通知", "请示", "函", "报告"],
                value="通知",
                label="公文类型"
            )
            doc_requirements = gr.Textbox(
                label="公文创作要求",
                placeholder="请输入详细的公文创作要求...",
                lines=30,
                interactive=True
            )
            doc_generate_btn = gr.Button("生成公文", variant="primary")
        with gr.Column(scale=3, elem_classes="large-output"):
            doc_history = gr.State([])
            doc_output = gr.Chatbot(
                value=[],
                label="公文生成结果",
                height="80vh",
                elem_classes="document-chat",
                type="messages"
            )
            with gr.Row():
                doc_download_btn = gr.Button("下载Word文档", visible=False)
                doc_preview = gr.Textbox(
                    label="公文预览",
                    lines=30,
                    interactive=False,
                    elem_classes="document-preview"
                )

    # 小说功能绑定
    download_button.click(
        download_novel,
        outputs=gr.File(label="下载小说")
    )
    gen_ouline_button.click(
        gen_ouline_button_clicked,
        [aign, user_idea_text, chatBox],
        [aign, chatBox, novel_outline_text, gen_ouline_button],
    )
    gen_beginning_button.click(
        gen_beginning_button_clicked,
        [
            aign,
            chatBox,
            novel_outline_text,
            user_requriments_text,
            embellishment_idea_text,
        ],
        [
            aign,
            chatBox,
            writing_plan_text,
            temp_setting_text,
            novel_content_text,
            gen_beginning_button,
        ],
    )
    gen_next_paragraph_button.click(
        gen_next_paragraph_button_clicked,
        [
            aign,
            chatBox,
            user_idea_text,
            novel_outline_text,
            writing_memory_text,
            temp_setting_text,
            writing_plan_text,
            user_requriments_text,
            embellishment_idea_text,
        ],
        [
            aign,
            chatBox,
            writing_plan_text,
            temp_setting_text,
            writing_memory_text,
            novel_content_text,
            gen_next_paragraph_button,
        ],
    )
    gen_ending_button.click(
        gen_ending_button_clicked,
        [
            aign,
            chatBox,
            user_requriments_text,
            novel_content_text
        ],
        [
            aign,
            chatBox,
            writing_plan_text,
            temp_setting_text,
            writing_memory_text,
            novel_content_text,
            gen_ending_button
        ]
    )

    # 公文功能绑定
    doc_generate_btn.click(
        fn=process_document_request,
        inputs=[doc_type, doc_requirements, doc_history],
        outputs=[doc_output, doc_preview, doc_download_btn],
        queue=True
    )
    doc_download_btn.click(
        fn=download_document,
        inputs=[doc_type, doc_preview],
        outputs=gr.File(label="下载公文")
    )

    # 界面切换
    switch_btn.click(
        lambda mode: "document" if mode == "fiction" else "fiction",
        inputs=[current_mode],
        outputs=[current_mode],
        queue=False
    ).then(
        toggle_interface,
        inputs=[current_mode],
        outputs=[fiction_interface, document_interface, switch_btn],
        queue=False
    )

demo.queue()
share_url = demo.launch(share=True)
print(f"分享此链接以允许他人访问您的应用: {share_url}")
