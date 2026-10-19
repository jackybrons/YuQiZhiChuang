import random
import threading
import time
import base64
import os
from dotenv import load_dotenv

load_dotenv()

import gradio as gr
from AIGC import GovernmentDocumentSystem
from AIGN import AIGN
from SCRIPT import ScriptAIGN
from zhipuAI import qwenChatLLM
import re
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import markdown
from bs4 import BeautifulSoup
from voice_input import record_audio, realtime_asr,signal_handler
import voice_input
import threading
import signal
import sounddevice as sd
import sys
import traceback
import queue
import soundfile as sf
import sounddevice as sd
import base64
import numpy as np
from scipy.signal import find_peaks
from aliyunsdkcore.client import AcsClient
from aliyunsdkcore.acs_exception.exceptions import ClientException, ServerException
from aliyunsdkcore.request import CommonRequest
import json
from tencentcloud.common import credential
from tencentcloud.common.profile.client_profile import ClientProfile
from tencentcloud.common.profile.http_profile import HttpProfile
from tencentcloud.asr.v20190614 import asr_client, models

# 自定义 excepthook
def custom_excepthook(type, value, tb):
    try:
        print("Custom excepthook triggered.")
        print(f"Exception type: {type.__name__}")
        print(f"Exception message: {value}")
        print("Traceback:")
        traceback.print_tb(tb)
    except Exception as e:
        print(f"Error in custom excepthook: {e}")
        
class Config:
    # 腾讯云配置
    MODEL = "qwen-plus-latest"
    SECRET_ID = os.environ.get("TENCENT_SECRET_ID", "")
    SECRET_KEY = os.environ.get("TENCENT_SECRET_KEY", "")
    REGION = "ap-guangzhou"

    # Mindcraft配置
    MINDCRAFT_API_KEY = os.environ.get("MINDCRAFT_API_KEY", "")
    MINDCRAFT_API_URL = "https://api.mindcraft.com.cn/v1/chat/completions"

    # 阿里云配置（更新为已验证的配置）
    ALIYUN_ACCESS_ID = os.environ.get("ALIYUN_ACCESS_KEY_ID", "")
    ALIYUN_ACCESS_KEY = os.environ.get("ALIYUN_ACCESS_KEY_SECRET", "")
    ALIYUN_REGION = "cn-hangzhou"
    ALIYUN_NLP_ENDPOINT = "alinlp.cn-hangzhou.aliyuncs.com"
    ALIYUN_NLP_VERSION = "2020-06-29"
    ALIYUN_NLP_SERVICE = "alinlp"

    # 音频参数
    RATE = 16000
    CHANNELS = 1
    DTYPE = 'int16'
    BUFFER_DURATION = 0.1
    WAVE_OUTPUT_FILENAME = "temp_record.wav"
    NORMALIZED_FILENAME = "normalized.wav"
    CHUNK_SIZE = 1024

chatLLM = qwenChatLLM()
STREAM_INTERVAL = 0.

# 公文生成常量
STANDARD_FONT = "SimSun"
FONT_MAPPING = {
    "SimSun": ["宋体", "SimSun"],
    "SimHei": ["黑体", "SimHei"],
    "FangSong": ["仿宋", "FangSong_GB2312", "FangSong"]
}

DOC_TYPES = {
    "通知": {
        "structure": ["标题", "发文字号", "主送机关", "正文", "发文机关", "成文日期"],
        "format_rules": {
            "title": {"font": "SimHei", "size": 22, "align": "center"},
            "body": {"font": STANDARD_FONT, "size": 16}
        }
    },
    "请示": {
        "structure": ["标题", "主送机关", "请示缘由", "请示事项", "结语"],
        "format_rules": {
            "title": {"font": "SimHei", "size": 22},
            "emphasis": {"bold": True, "underline": True}
        }
    },
    "函": {
        "structure": ["标题", "主送机关", "发函缘由", "商洽事项", "结语"],
        "format_rules": {
            "title": {"font": "SimHei", "size": 20}
        }
    },
    "报告": {
        "structure": ["标题", "主送机关", "基本情况", "分析情况", "意见建议"],
        "format_rules": {
            "title": {"font": "SimHei", "size": 20}
        }
    }
}

audio_queue = queue.Queue()
is_recording = threading.Event()
stop_recording = threading.Event()
current_recording_target = None

def print_color(text, color='green'):
    colors = {
        'red': '\033[91m',
        'green': '\033[92m',
        'yellow': '\033[93m',
        'blue': '\033[94m',
        'end': '\033[0m'
    }
    print(f"{colors[color]}{text}{colors['end']}")

def record_audio():
    """录制音频"""
    print_color("\n🔴 正在录音... (按Ctrl+C停止)", 'red')
    try:
        with sf.SoundFile(Config.WAVE_OUTPUT_FILENAME, mode='w',
                          samplerate=Config.RATE, channels=Config.CHANNELS) as f:
            with sd.InputStream(samplerate=Config.RATE,
                                channels=Config.CHANNELS,
                                callback=audio_callback):
                while not stop_recording.is_set():
                    if not audio_queue.empty():
                        data = audio_queue.get()
                        if check_audio_quality(data, Config.RATE):
                            f.write(data)
            print("录音线程检测到停止标志，正在退出...")
        stop_recording.clear()  # 录音结束后清除停止标志
        return Config.WAVE_OUTPUT_FILENAME
    except KeyboardInterrupt:
        print_color("\n检测到Ctrl+C，停止录音...", 'yellow')
        stop_recording.set()
        while is_recording.is_set():
            time.sleep(0.1)
        return Config.WAVE_OUTPUT_FILENAME
    except Exception as e:
        print_color(f"录音失败: {str(e)}", 'red')
        return None

def audio_callback(indata, frames, time, status):
    """音频回调函数"""
    if status:
        print_color(f"音频流状态: {status}", 'yellow')
    audio_queue.put(indata.copy())

def check_audio_quality(data, rate):
    """简单音频质量检查"""
    if np.max(np.abs(data)) < 0.00001:
        print_color("警告: 检测到静音或音量过低", 'yellow')
        return False
    return True

def init_tencent_client():
    """初始化腾讯云客户端"""
    cred = credential.Credential(Config.SECRET_ID, Config.SECRET_KEY)
    http_profile = HttpProfile()
    http_profile.endpoint = "asr.tencentcloudapi.com"

    client_profile = ClientProfile()
    client_profile.httpProfile = http_profile

    return asr_client.AsrClient(cred, Config.REGION, client_profile)

def realtime_asr(audio_file):
    """调用腾讯云ASR进行语音识别"""
    try:
        with open(audio_file, 'rb') as f:
            audio_data = f.read()

        client = init_tencent_client()
        req = models.SentenceRecognitionRequest()

        req.ProjectId = 0
        req.SubServiceType = 2
        req.EngSerViceType = "16k_zh"
        req.SourceType = 1
        req.Data = base64.b64encode(audio_data).decode()
        req.DataLen = len(audio_data)
        req.VoiceFormat = "wav"

        resp = client.SentenceRecognition(req)
        if resp.Result:
            return resp.Result
        return None
    except Exception as e:
        print_color(f"语音识别失败: {str(e)}", 'red')
        return None

def handle_voice_input(target_id):
    """开始录音并设置当前录音目标"""
    global is_recording, stop_recording, current_recording_target
    
    # 设置当前录音目标
    current_recording_target = target_id
    
    try:
        if not is_recording.is_set():
            print(f"开始录音，目标输入框: {target_id}")
            is_recording.set()
            stop_recording.clear()
            
            # 开始录音线程
            audio_thread = threading.Thread(target=record_audio)
            audio_thread.start()
            
            # 根据目标ID找到对应的输入框组件
            target_component = None
            if target_id == "fiction_idea":
                target_component = user_idea_text
            elif target_id == "fiction_requirements":
                target_component = user_requriments_text
            elif target_id == "fiction_embellishment":
                target_component = embellishment_idea_text
            elif target_id == "document_requirements":
                target_component = doc_requirements
            elif target_id == "script_idea":
                target_component = script_user_idea
            elif target_id == "script_requirements":
                target_component = script_user_requirements
            
            if target_component:
                return gr.update(value="(录音中...)")
            
    except Exception as e:
        print(f"录音启动异常: {e}")
    
    return gr.update()

# 2. 修改stop_voice_recording函数，添加到mix_all1.py中（替换原有函数）
def stop_voice_recording(target_id):
    """停止录音并将识别结果输出到对应输入框"""
    global is_recording, stop_recording, current_recording_target
    
    # 验证目标ID是否匹配
    if target_id != current_recording_target:
        print(f"警告: 停止录音的目标ID ({target_id}) 与当前记录的目标ID ({current_recording_target}) 不匹配")
        return gr.update()
    
    try:
        if is_recording.is_set():
            print(f"停止录音，目标输入框: {target_id}")
            stop_recording.set()
            
            # 等待录音线程停止
            max_wait = 3
            start_time = time.time()
            while is_recording.is_set() and (time.time() - start_time) < max_wait:
                time.sleep(0.1)
            
            # 强制停止音频设备
            sd.stop()
            is_recording.clear()
            stop_recording.clear()
            
            # 获取实际的录音文件路径
            audio_file = Config.WAVE_OUTPUT_FILENAME
            
            # 语音识别
            text = ""
            if os.path.exists(audio_file) and os.path.getsize(audio_file) > 0:
                try:
                    text = realtime_asr(audio_file)
                    if not text:
                        text = "(未识别到语音内容)"
                except Exception as e:
                    print(f"语音识别异常: {e}")
                    text = "(语音识别失败)"
                
                # 清理临时文件
                try:
                    os.remove(audio_file)
                except:
                    pass
            else:
                text = "(未检测到录音文件)"
            
            print(f"语音识别结果: {text}")
            
            # 根据目标ID返回结果到对应输入框
            if target_id == "fiction_idea":
                return gr.update(value=text)
            elif target_id == "fiction_requirements":
                return gr.update(value=text)
            elif target_id == "fiction_embellishment":
                return gr.update(value=text)
            elif target_id == "document_requirements":
                return gr.update(value=text)
            elif target_id == "script_idea":
                return gr.update(value=text)
            elif target_id == "script_requirements":
                return gr.update(value=text)
    
    except Exception as e:
        print(f"停止录音异常: {e}")
    
    return gr.update()

def record_audio():
    """录制音频"""
    print_color("\n🔴 正在录音... (按Ctrl+C停止)", 'red')
    
    try:
        with sf.SoundFile(Config.WAVE_OUTPUT_FILENAME, mode='w',
                         samplerate=Config.RATE, channels=Config.CHANNELS) as f:
            with sd.InputStream(samplerate=Config.RATE,
                               channels=Config.CHANNELS,
                               callback=audio_callback):
                while not stop_recording.is_set():
                    if not audio_queue.empty():
                        data = audio_queue.get()
                        if check_audio_quality(data, Config.RATE):
                            f.write(data)
            print("录音线程检测到停止标志，正在退出...")  # 添加调试信息
        stop_recording.clear()  # 录音结束后清除停止标志
        return Config.WAVE_OUTPUT_FILENAME
    except KeyboardInterrupt:
        print_color("\n检测到Ctrl+C，停止录音...", 'yellow')
        stop_recording.set()
        while is_recording.is_set():
            time.sleep(0.1)
        return Config.WAVE_OUTPUT_FILENAME
    except Exception as e:
        print_color(f"录音失败: {str(e)}", 'red')
        return None

def signal_handler(sig, frame):
    global is_recording, stop_recording
    if is_recording.is_set():
        stop_recording.set()
        print("停止录音...")
        # 等待录音结束
        while is_recording.is_set():
            time.sleep(0.1)
    else:
        print("程序正常退出...")
        exit(0)

# 注册信号处理函数
signal.signal(signal.SIGINT, signal_handler)

def convert_history(history):
    converted = []
    for msg in history:
        if isinstance(msg, (list, tuple)) and len(msg) == 2:
            converted.append(msg)
        elif isinstance(msg, dict):
            content = msg.get("content", "")
            if msg.get("role") == "user":
                converted.append([None, content])
            elif msg.get("role") == "assistant":
                converted.append([content, None])
    return converted

def make_middle_chat():
    carrier = threading.Event()
    carrier.history = []
    def middle_chat(messages, temperature=None, top_p=None):
        nonlocal carrier
        carrier.history.append([None, ""])
        if len(carrier.history) > 20:
            carrier.history = carrier.history[-16:]
        try:
            output_text = ''
            for resp in chatLLM(messages, temperature=temperature, top_p=top_p, stream=True):
                output_text += resp["content"]
                carrier.history[-1] = [output_text, None]
            return {"content": output_text}
        except Exception as e:
            carrier.history[-1] = [f"Error: {e}", None]
            raise e
    return carrier, middle_chat

def get_available_font(font_name: str) -> str:
    for font in FONT_MAPPING.get(font_name, [font_name]):
        return font
    return "Arial Unicode MS"

def sanitize_text(text: str) -> str:
    text = re.sub(r'[\u3000\xa0\u200b]', ' ', text)
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
    return text.strip()

def format_chinese_date():
    return datetime.now().strftime("%Y年%m月%d日")

def _generate_doc_number(doc_type: str) -> str:
    prefix_map = {"通知": "X政", "请示": "X政请", "函": "X政函", "报告": "X政报"}
    year = datetime.now().year
    seq = f"{int(time.time()) % 1000:03d}"
    return f"{prefix_map.get(doc_type, 'X政')}〔{year}〕{seq}号"

def _set_font(element, font_name: str):
    rPr = element._element.rPr
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        element._element.insert(0, rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)

def generate_document(doc_type, sender, recipient, content, doc_number, doc_date):
    try:
        full_input = {
            "type": doc_type,
            "sender": sender,
            "recipient": recipient,
            "content": content,
            "doc_number": doc_number or _generate_doc_number(doc_type),
            "date": doc_date or format_chinese_date()
        }
        
        messages = [
            {"role": "system", "content": f"请按标准格式生成{doc_type}，直接返回最终公文内容"},
            {"role": "user", "content": content}
        ]
        
        print(f"请求的 messages: {messages}")
        response = chatLLM(messages)
        print(f"响应的 response: {response}")
        result = response["content"]
        
        lines = result.split('\n')
        formatted = []
        for i, line in enumerate(lines):
            if i in [0, 1, 3] or i >= len(lines)-2:
                formatted.append(line)
            else:
                if line.strip():
                    formatted.append(f"　　{line}" if not line.startswith(' ') else line)
        
        return "\n".join(formatted)
    except Exception as e:
        print(f"公文生成失败，详细错误信息: {traceback.format_exc()}")
        return f"生成失败: {str(e)}"
    
def export_to_word(content_text, doc_type):
    try:
        lines = content_text.split('\n')
        doc_info = {
            "type": doc_type,
            "title": lines[0] if lines else "",
            "doc_number": lines[1] if len(lines)>1 else "",
            "recipient": lines[3] if len(lines)>3 else "",
            "content": '\n'.join(lines[4:-2]) if len(lines)>5 else "",
            "sender": lines[-2] if len(lines)>=2 else "",
            "date": lines[-1] if lines else ""
        }
        
        output_path = os.path.join("output", f"{doc_type}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx")
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        
        document = Document()
        section = document.sections[0]
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.5)
        
        # 标题
        title_para = document.add_paragraph()
        title_run = title_para.add_run(sanitize_text(doc_info["title"]))
        _set_font(title_run, get_available_font("SimHei"))
        title_run.font.size = Pt(22)
        title_run.font.bold = True
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph()
        
        # 文号
        if doc_info.get("doc_number"):
            num_para = document.add_paragraph()
            num_run = num_para.add_run(sanitize_text(doc_info["doc_number"]))
            _set_font(num_run, get_available_font(STANDARD_FONT))
            num_run.font.size = Pt(16)
            num_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            document.add_paragraph()
        
        # 正文内容
        body_font = get_available_font(STANDARD_FONT)
        body_size = Pt(16)
        
        content = sanitize_text(doc_info["content"])
        for paragraph in content.split('\n'):
            if not paragraph.strip():
                document.add_paragraph()
                continue
                
            p = document.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(32)
            p.paragraph_format.line_spacing = 1.5
            
            run = p.add_run(paragraph)
            _set_font(run, body_font)
            run.font.size = body_size

        # 发文信息
        if all(key in doc_info for key in ["sender", "date"]):
            document.add_paragraph()
            
            sender_para = document.add_paragraph()
            sender_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            sender_run = sender_para.add_run(sanitize_text(doc_info["sender"]))
            _set_font(sender_run, body_font)
            sender_run.font.size = body_size
            
            date_para = document.add_paragraph()
            date_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            date_run = date_para.add_run(sanitize_text(doc_info["date"]))
            _set_font(date_run, body_font)
            date_run.font.size = body_size

        document.save(output_path)
        return output_path
    except Exception as e:
        raise gr.Error(f"Word导出失败: {str(e)}")

def gen_ouline_button_clicked(aign, user_idea, history, chapter_count, word_count_per_chapter):
    history = convert_history(history)
    aign.user_idea = user_idea
    aign.chapter_count = chapter_count
    aign.word_count_per_chapter = word_count_per_chapter

    carrier, middle_chat = make_middle_chat()
    carrier.history = history
    aign.novel_outline_writer.chatLLM = middle_chat

    gen_ouline_thread = threading.Thread(target=aign.genNovelOutline)
    gen_ouline_thread.start()

    while gen_ouline_thread.is_alive():
        yield [
            aign,
            carrier.history,
            aign.novel_outline,
            gr.Button(visible=True),
            gr.Button(visible=False),
            gr.Button(visible=False),
            gr.Button(visible=False)
        ]
        time.sleep(STREAM_INTERVAL)
        
    yield [
        aign,
        carrier.history,
        aign.novel_outline,
        gr.Button(visible=True),
        gr.Button(visible=True),
        gr.Button(visible=False),
        gr.Button(visible=False)
    ]

def gen_beginning_button_clicked(aign, history, novel_outline, user_requriments, embellishment_idea):
    history = convert_history(history)
    aign.novel_outline = novel_outline
    aign.user_requriments = user_requriments
    aign.embellishment_idea = embellishment_idea

    # 将大纲添加到对话历史时分割成长度合适的段落
    outline_parts = [novel_outline[i:i+500] for i in range(0, len(novel_outline), 500)]
    updated_history = history.copy()
    
    for part in outline_parts:
        updated_history.append([None, f"大纲片段:\n{part}"])
    
    carrier, middle_chat = make_middle_chat()
    carrier.history = updated_history
    aign.novel_beginning_writer.chatLLM = middle_chat
    aign.novel_embellisher.chatLLM = middle_chat

    gen_beginning_thread = threading.Thread(target=aign.genBeginning)
    gen_beginning_thread.start()

    while gen_beginning_thread.is_alive():
        # 实时更新时也分割长文本
        if aign.novel_content and len(aign.novel_content) > 500:
            content_parts = [aign.novel_content[i:i+500] for i in range(0, len(aign.novel_content), 500)]
            temp_history = carrier.history + [(None, f"生成中...\n{part}") for part in content_parts]
        else:
            temp_history = carrier.history + [(None, "生成中..." + (aign.novel_content[:500] + "..." if aign.novel_content else ""))]
        
        yield [
            aign,
            temp_history,
            aign.writing_plan,
            aign.temp_setting,
            aign.novel_content,
            gr.Button(visible=True),
            gr.Button(visible=True),
            gr.Button(visible=False)
        ]
        time.sleep(STREAM_INTERVAL)
    
    # 最终输出时也分割长文本
    final_content = aign.novel_content
    if len(final_content) > 1000:
        final_parts = [final_content[i:i+500] for i in range(0, len(final_content), 500)]
        final_history = carrier.history + [(None, f"最终生成:\n{part}") for part in final_parts]
    else:
        final_history = carrier.history + [(None, f"最终生成:\n{final_content}")]
    
    yield [
        aign,
        final_history,
        aign.writing_plan,
        aign.temp_setting,
        aign.novel_content,
        gr.Button(visible=True),
        gr.Button(visible=True),
        gr.Button(visible=True)
    ]

def gen_next_paragraph_button_clicked(
    aign,
    history,
    user_idea,
    novel_outline,
    writing_memory,
    temp_setting,
    writing_plan,
    user_requriments,
    embellishment_idea,
    auto_gen_next
):
    history = convert_history(history)
    current_state = {
        "paragraph_list": aign.paragraph_list.copy(),
        "writing_plan": aign.writing_plan,
        "temp_setting": aign.temp_setting,
        "writing_memory": aign.writing_memory,
        "novel_content": aign.novel_content
    }
    aign.history_stack.append(current_state)
    if len(aign.history_stack) > aign.max_history:
        aign.history_stack.pop(0)
        
    aign.user_idea = user_idea
    aign.novel_outline = novel_outline
    aign.writing_memory = writing_memory
    aign.temp_setting = temp_setting
    aign.writing_plan = writing_plan
    aign.user_requriments = user_requriments
    aign.embellishment_idea = embellishment_idea

    carrier, middle_chat = make_middle_chat()
    carrier.history = history
    aign.novel_writer.chatLLM = middle_chat
    aign.novel_embellisher.chatLLM = middle_chat
    aign.memory_maker.chatLLM = middle_chat

    # 获取当前章节信息
    current_chapter = len(aign.paragraph_list) + 1
    total_chapters = aign.chapter_count
    
    time.sleep(0.2)
    
    gen_next_paragraph_thread = threading.Thread(target=aign.genNextParagraph)
    gen_next_paragraph_thread.start()

    while gen_next_paragraph_thread.is_alive():
        yield [
            aign,
            carrier.history,
            aign.writing_plan,
            aign.temp_setting,
            aign.writing_memory,
            aign.novel_content,
            gr.Button(visible=True),
            gr.Button(visible=True),
            gr.Button(visible=True)
        ]
        time.sleep(STREAM_INTERVAL)
        
    # 添加章节标记
    if aign.novel_content and f"第{current_chapter}章" not in aign.novel_content:
        chapter_header = f"\n\n# 第{current_chapter}章/{total_chapters}章\n\n"
        aign.novel_content = aign.novel_content + chapter_header + aign.paragraph_list[-1]
    
    yield [
        aign,
        carrier.history,
        aign.writing_plan,
        aign.temp_setting,
        aign.writing_memory,
        aign.novel_content,
        gr.Button(visible=True),
        gr.Button(visible=True),
        gr.Button(visible=True)
    ]

    final_history = carrier.history + [
            (None, f"更新后的写作计划:\n{aign.writing_plan}"),
            (None, f"更新后的临时设定:\n{aign.temp_setting}"),
            (None, f"新增段落:\n{aign.paragraph_list[-1] if aign.paragraph_list else ''}")
        ]
        
    yield [
            aign,
            final_history,  # 使用整合后的历史
            aign.writing_plan,
            aign.temp_setting,
            aign.writing_memory,
            aign.novel_content,
            gr.Button(visible=True),
            gr.Button(visible=True),
            gr.Button(visible=True)
        ]
def gen_ending_button_clicked(aign, history, user_requriments, novel_content):
    history = convert_history(history)
    aign.user_requriments = user_requriments
    
    carrier, middle_chat = make_middle_chat()
    carrier.history = history
    aign.novel_ending_writer.chatLLM = middle_chat
    aign.novel_embellisher.chatLLM = middle_chat

    current_state = {
        "paragraph_list": aign.paragraph_list.copy(),
        "writing_plan": aign.writing_plan,
        "temp_setting": aign.temp_setting,
        "writing_memory": aign.writing_memory,
        "novel_content": aign.novel_content
    }
    aign.history_stack.append(current_state)
    if len(aign.history_stack) > aign.max_history:
        aign.history_stack.pop(0)

    gen_ending_thread = threading.Thread(target=aign.genEnding)
    gen_ending_thread.start()

    while gen_ending_thread.is_alive():
        yield [
            aign,
            carrier.history,
            aign.writing_plan,
            aign.temp_setting,
            aign.writing_memory,
            aign.novel_content,
            gr.Button(visible=True),
            gr.Button(visible=True),
            gr.Button(visible=True)
        ]
        time.sleep(STREAM_INTERVAL)
        
    yield [
        aign,
        carrier.history,
        aign.writing_plan,
        aign.temp_setting,
        aign.writing_memory,
        aign.novel_content,
        gr.Button(visible=True),
        gr.Button(visible=True),
        gr.Button(visible=True)
    ]

def download_novel():
    if os.path.exists("novel_record.md"):
        with open("novel_record.md", "r+", encoding="utf-8") as f:
            content = f.read()
            if "[END]" not in content:
                f.write("\n\n# [END]\n")
        return "novel_record.md"
    raise gr.Error("尚未生成任何小说内容")

def download_document(doc_type, content):
    try:
        if not content.strip():
            raise gr.Error("没有可下载的内容")
        
        output_path = export_to_word(content, doc_type)
        if os.path.exists(output_path):
            return output_path
        raise gr.Error("文件生成失败")
    except Exception as e:
        raise gr.Error(f"下载失败: {str(e)}")

def toggle_interface(mode):
    is_fiction = mode == "fiction"
    is_script = mode == "script"
    return {
        fiction_interface: gr.Row(visible=is_fiction),
        document_interface: gr.Row(visible=not (is_fiction or is_script)),
        script_interface: gr.Row(visible=is_script),
        switch_btn: gr.Button(
            value="切换到公文模式" if is_fiction else ("切换到剧本模式" if mode == "document" else "切换到小说模式"),
            variant="primary" if not is_fiction else "secondary"
        )
    }

def process_document_request(doc_type, doc_requirements, history):
    history = convert_history(history)
    
    try:
        sender = "XX市城市绿化管理局"
        recipient = "各区县城市绿化管理部门、各相关单位"
        doc_number = _generate_doc_number(doc_type)
        doc_date = format_chinese_date()
        
        generated_content = generate_document(
            doc_type=doc_type,
            sender=sender,
            recipient=recipient,
            content=doc_requirements,
            doc_number=doc_number,
            doc_date=doc_date
        )
        
        # 修改这部分，将生成的内容同时添加到对话历史中
        user_message = f"请帮我生成一份{doc_type}，具体要求如下：\n{doc_requirements}"
        assistant_message = f"📝 已为您生成{doc_type}，内容如下：\n\n{generated_content}"
        
        # 更新对话历史
        updated_history = history + [
            (user_message, assistant_message)
        ]
        
        return [
            updated_history,  # 更新后的对话历史
            generated_content,  # 公文预览内容
            gr.Button(visible=True)  # 显示下载按钮
        ]
    
    except Exception as e:
        error_msg = f"❌ {doc_type}生成失败：{str(e)}"
        return [
            history + [(doc_requirements, error_msg)],  # 错误信息添加到对话历史
            "",  # 清空预览
            gr.Button(visible=False)  # 隐藏下载按钮
        ]
        
def _generate_doc_number(doc_type: str) -> str:
    prefix_map = {"通知": "X政", "请示": "X政请", "函": "X政函", "报告": "X政报"}
    year = datetime.now().year
    seq = f"{int(time.time()) % 1000:03d}"
    return f"{prefix_map.get(doc_type, 'X政')}〔{year}〕{seq}号"

def format_chinese_date():
    return datetime.now().strftime("%Y年%m月%d日")

def _set_font(element, font_name: str):
    rPr = element._element.rPr
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        element._element.insert(0, rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)

def get_available_font(font_name: str) -> str:
    for font in FONT_MAPPING.get(font_name, [font_name]):
        return font
    return "Arial Unicode MS"

def sanitize_text(text: str) -> str:
    text = re.sub(r'[\u3000\xa0\u200b]', ' ', text)
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
    return text.strip()

# 剧本生成相关函数
def gen_script_outline_button_clicked(aign, user_idea, history, scene_count, word_count_per_scene):
    history = convert_history(history)
    aign.user_idea = user_idea
    aign.scene_count = scene_count
    aign.word_count_per_scene = word_count_per_scene

    carrier, middle_chat = make_middle_chat()
    carrier.history = history
    aign.creative_architect.chatLLM = middle_chat

    thread = threading.Thread(target=lambda: aign.genScriptOutline())
    thread.start()
    while thread.is_alive():
        yield [aign, carrier.history, aign.script_outline, gr.update(visible=True)]
        time.sleep(STREAM_INTERVAL)
    yield [aign, carrier.history, aign.script_outline, gr.update(visible=True)]

def gen_script_beginning_button_clicked(aign, history, script_outline, user_requirements):
    history = convert_history(history)
    aign.script_outline = script_outline
    aign.user_requirements = user_requirements
    # 初始化场景计数器
    aign.total_scenes = aign.scene_count  
    aign.current_scene = 0
    
    
    carrier, middle_chat = make_middle_chat()
    carrier.history = history
    aign.master_director.chatLLM = middle_chat
    aign.scene_engineer.chatLLM = middle_chat
    # 仅生成开头
    thread = threading.Thread(target=aign.genBeginning)
    thread.start()
    
    while thread.is_alive():
        current_content = aign.scene_list[0] if aign.scene_list else ""
        yield [aign, carrier.history, aign.working_plan, aign.temp_setting, aign.script_content, gr.update(visible=True)]
        time.sleep(STREAM_INTERVAL)
    yield [aign, carrier.history, aign.working_plan, aign.temp_setting, aign.script_content, current_content, gr.Button(visible=True)]

def gen_next_scene_button_clicked(
    aign,
    history,
    user_idea,
    script_outline,
    script_memory,
    temp_setting,
    working_plan,
    user_requirements,
    auto_gen_next
):

    history = convert_history(history)
    if len(aign.history_stack) >= aign.max_history:
        aign.history_stack.pop(0)
    aign.history_stack.append({
        "scene_list": aign.scene_list.copy(),
        "script_memory": aign.script_memory,
        "script_content": aign.script_content
    })

    aign.user_idea = user_idea
    aign.script_outline = script_outline
    aign.script_memory = script_memory
    aign.temp_setting = temp_setting
    aign.working_plan = working_plan
    aign.user_requirements = user_requirements

    carrier, middle_chat = make_middle_chat()
    carrier.history = history
    aign.master_director.chatLLM = middle_chat
    aign.dialogue_alchemist.chatLLM = middle_chat
    aign.cultural_guardian.chatLLM = middle_chat

    # 获取当前场景信息
    current_scene = len(aign.scene_list) + 1
    total_scenes = aign.scene_count
    
    # 分步生成
    thread = threading.Thread(target=aign.genNextScene)
    thread.start()
    while thread.is_alive():
        progress = f"当前进度：{aign.current_scene}/{aign.total_scenes}场"
        yield [aign, carrier.history, aign.working_plan, aign.temp_setting, aign.script_memory, aign.script_content, progress + "\n" + aign.scene_list[-1]]
        time.sleep(STREAM_INTERVAL)
    
    # 添加场景标记
    if aign.script_content and f"第{current_scene}场" not in aign.script_content:
        scene_header = f"\n\n## 第{current_scene}场/{total_scenes}场\n\n"
        aign.script_content = aign.script_content + scene_header + aign.scene_list[-1]
    
    yield [aign, carrier.history, aign.working_plan, aign.temp_setting, aign.script_memory, aign.script_content]

def gen_script_ending_button_clicked(aign, history, user_requirements, script_content):
    history = convert_history(history)
    aign.user_requirements = user_requirements

    carrier, middle_chat = make_middle_chat()
    carrier.history = history
    aign.master_director.chatLLM = middle_chat
    aign.dialogue_alchemist.chatLLM = middle_chat

    if len(aign.history_stack) >= aign.max_history:
        aign.history_stack.pop(0)
    aign.history_stack.append({
        "scene_list": aign.scene_list.copy(),
        "script_memory": aign.script_memory,
        "script_content": aign.script_content
    })

    try:
        thread = threading.Thread(target=lambda: aign.genEnding())
        thread.start()
        while thread.is_alive():
            yield [aign, carrier.history, aign.working_plan, aign.temp_setting, aign.script_memory, aign.script_content, gr.update(visible=True)]
            time.sleep(STREAM_INTERVAL)
        yield [aign, carrier.history, aign.working_plan, aign.temp_setting, aign.script_memory, aign.script_content, gr.update(visible=True)]
    except Exception as e:
        aign.restore_previous_state()
        raise gr.Error(f"结局生成失败，已回滚: {e}")

def download_script():
    if os.path.exists("script_record.md"):
        with open("script_record.md", "r+", encoding="utf-8") as f:
            content = f.read()
            if "[END]" not in content:
                f.write("\n\n# [END]\n")
        return "script_record.md"
    raise gr.Error("尚未生成任何剧本内容")

css = """
/* 主容器透明化 */
gradio-app {
    background-color: rgba(0,0,0,0) !important;
    background-image: none !important;
}

/* 全局文字颜色 */
* {
    color: black !important;
}

/* 更精确的全局背景设置 */
body {
    background: #e6f7ff !important;
}

/* 新增大尺寸输入输出框样式 */
.large-input .gr-input,
.large-output .gr-textbox {
    min-height: 600px !important;
    font-size: 14px !important;
    line-height: 1.6 !important;
}

/* 标题行样式 */
.title-row {
    margin-bottom: 20px !important;
}

.mode-button {
    margin-left: 20px !important;
    background: #2196F3 !important;
    color: white !important;
    padding: 6px 12px !important;
    font-size: 14px !important;
    border-radius: 4px !important;
    border: none !important;
    cursor: pointer !important;
}

.document-chat .message-bot {
    font-family: 'SimSun', serif !important;
    white-space: pre-wrap;
    background: #f0faff !important;
    border-left: 4px solid #1890ff !important;
}

.document-preview {
    font-family: 'FangSong', serif !important;
    font-size: 16px !important;
    line-height: 2 !important;
    padding: 20px !important;
    background: #f9f9f9 !important;
    border: 1px solid #ddd !important;
}

/* 完全重写下拉菜单样式 */
.document-interface .gr-dropdown {
    --tw-bg-opacity: 1 !important;
    background-color: rgb(255 255 255 / var(--tw-bg-opacity)) !important;
}

.document-interface .gr-dropdown .wrap {
    background-color: white !important;
    background-image: none !important;
    border-radius: 4px !important;
    border: 1px solid #ddd !important;
}

.document-interface .gr-dropdown .wrap:hover {
    background-color: #f5f5f5 !important;
}

.document-interface .gr-dropdown .options {
    background-color: white !important;
    background-image: none !important;
    border: 1px solid #ddd !important;
    box-shadow: 0 2px 4px rgba(0,0,0,0.1) !important;
}

.document-interface .gr-dropdown .option {
    background-color: white !important;
    background-image: none !important;
}

.document-interface .gr-dropdown .option:hover {
    background-color: #f0f7ff !important;
}

/* 添加对Gradio 3.x+的特殊处理 */
.document-interface .gr-box {
    background-color: white !important;
    background-image: none !important;
}

/* 保持原有其他样式不变 */
.document-preview {
    font-family: 'FangSong', serif !important;
    font-size: 16px !important;
    line-height: 2 !important;
    padding: 20px !important;
    background: #f9f9f9 !important;
    border: 1px solid #ddd !important;
}
.hidden {
    display: none !important;
}

/* 剧本界面特定样式 */
.script-panel {
    background: rgba(255,253,245,0.8) !important;
    border: 1px solid #f5e8c4 !important;
}

.script-content {
    font-family: 'KaiTi', serif !important;
    font-size: 15px !important;
    line-height: 1.8 !important;
}

/* 章节设置样式 */
.chapter-settings {
    background: #f5f5f5 !important;
    padding: 15px !important;
    border-radius: 8px !important;
    margin-bottom: 15px !important;
}

.chapter-settings label {
    font-weight: bold !important;
    margin-bottom: 5px !important;
    display: block !important;
}

.chapter-settings input {
    width: 100% !important;
    padding: 8px !important;
    border: 1px solid #ddd !important;
    border-radius: 4px !important;
}

/* 章节标题样式 */
.chapter-title {
    font-weight: bold;
    color: #1890ff;
    margin: 15px 0 10px 0;
}

.scene-title {
    font-weight: bold;
    color: #722ed1;
    margin: 12px 0 8px 0;
}

/* 新增横向滑轮样式 */
.scrollable-chat {
    overflow-x: auto;
    white-space: nowrap;
    padding-bottom: 15px;
}

.scrollable-chat .message {
    display: inline-block;
    vertical-align: top;
    margin-right: 15px;
    white-space: normal;
    width: auto; /* 调整为自动宽度 */
    max-width: 700px; /* 设置最大宽度 */
    min-height: 100px;
    background: #f5f5f5;
    border-radius: 8px;
    padding: 10px;
    box-shadow: 0 2px 4px rgba(0,0,0,0.1);
    word-wrap: break-word; /* 允许单词换行 */
    overflow-wrap: break-word; /* 允许单词换行 */
}

.scrollable-chat .message-user {
    background: #e1f5fe;
    margin-left: auto; /* 修改 */
    margin-right: 0; /* 新增 */
    border-bottom-right-radius: 4px;
    border-left: 4px solid #4CAF50;
}

.scrollable-chat .message-bot {
    background: #f0f7ff;
    margin-left: 0; /* 修改 */
    margin-right: auto; /* 新增 */
    border-bottom-left-radius: 4px;
    border-left: 4px solid #2196F3;
}

/* 消息内容样式 */
.scrollable-chat .message-content {
    white-space: pre-wrap;
    overflow-wrap: break-word;
    word-break: break-word;
    max-width: 100%;
    font-size: 14px;
    line-height: 1.6;
}

/* 确保聊天容器正确处理溢出 */
.scrollable-chat {
    overflow-x: hidden !important;
    width: 100% !important;
}

/* 特别处理长文本消息 */
.scrollable-chat .message-bot.long-text {
    max-width: 90% !important;
    display: inline-block !important;
}

/* custom.css 或 style模块内 */
.novel-bubble-container {
    max-width: 100% !important;
    overflow-x: hidden !important;
    padding: 0 10px !important;
}

.novel-bubble-container .message-bot {
    max-width: calc(100% - 20px) !important;
    margin-left: 5px !important;
    background: #f5f7fa !important;
    border-radius: 12px !important;
    word-break: break-word !important;
    white-space: pre-wrap !important;
    overflow-wrap: anywhere !important;
    box-shadow: 0 1px 2px rgba(0,0,0,0.1) !important;
}

.novel-bubble-container .message-bot p {
    margin: 0.5em 0 !important;
}

"""

with gr.Blocks(css=css) as demo:
    aign = gr.State(AIGN(chatLLM))
    script_aign = gr.State(ScriptAIGN(chatLLM))
    current_mode = gr.State("fiction")
    
    with gr.Row(equal_height=True, elem_classes="title-row"):
        gr.Markdown("## 语启智创--文本创作助手")
        switch_btn = gr.Button("切换到公文模式", elem_classes="mode-button")
    
    # 小说创作界面
    with gr.Row(visible=True) as fiction_interface:
        with gr.Column(scale=1, elem_id="row1"):
            with gr.Tab("开始"):
                gr.Markdown("生成大纲->生成开头->生成下一段")
                user_idea_text = gr.Textbox(
                    placeholder="请在这里输入您的想法",
                    label="想法",
                    lines=7,
                    interactive=True,
                    elem_id="fiction_idea"
                )
                stop_voice_button_1 = gr.Button(value="录音完毕")
                stop_voice_button_1.click(
                    stop_voice_recording,
                    inputs=[gr.Textbox(value="fiction_idea", visible=False)],
                    outputs=[user_idea_text]
                )
                voice_input_button = gr.Button(value="语音输入")
                voice_input_button.click(
                    handle_voice_input,
                    inputs=[gr.Textbox(value="fiction_idea", visible=False)],
                    outputs=[user_idea_text]
                    )

                with gr.Accordion("章节设置", open=False, elem_classes="chapter-settings"):
                    gr.Markdown("**章节设置说明**：\n- 每章字数建议不少于500字\n- 章节数量建议5-30章")
                    chapter_count = gr.Number(
                        label="章节数量",
                        value=10,
                        minimum=1,
                        maximum=100,
                        step=1,
                        interactive=True
                    )
                    word_count_per_chapter = gr.Number(
                        label="每章节字数",
                        value=2000,
                        minimum=500,
                        maximum=10000,
                        step=100,
                        interactive=True
                    )
                user_requriments_text = gr.Textbox(
                    "",
                    label="写作要求",
                    lines=7,
                    interactive=True,
                    elem_id="fiction_requirements"
                )
                stop_voice_button_2 = gr.Button(value="录音完毕")
                stop_voice_button_2.click(
                    stop_voice_recording,
                    inputs=[gr.Textbox(value="fiction_requirements", visible=False)],
                    outputs=[user_requriments_text]
                )
                voice_input_button = gr.Button(value="语音输入")
                voice_input_button.click(
                    handle_voice_input,
                    inputs=[gr.Textbox(value="fiction_requirements", visible=False)],
                    outputs=[user_requriments_text]
                )

                embellishment_idea_text = gr.Textbox(
                    "",
                    label="润色要求",
                    lines=7,
                    interactive=True,
                    elem_id="fiction_embellishment"
                )
                stop_voice_button_3 = gr.Button(value="录音完毕")
                stop_voice_button_3.click(
                    stop_voice_recording,
                    inputs=[gr.Textbox(value="fiction_embellishment", visible=False)],
                    outputs=[embellishment_idea_text]
                )
                voice_input_button = gr.Button(value="语音输入")
                voice_input_button.click(
                    handle_voice_input,
                    inputs=[gr.Textbox(value="fiction_embellishment", visible=False)],
                    outputs=[embellishment_idea_text]
                )

                gen_ouline_button = gr.Button("生成大纲", variant="primary")
                
            with gr.Tab("大纲"):
                novel_outline_text = gr.Textbox(
                    "",
                    label="小说大纲",
                    lines=15,
                    interactive=True,
                    elem_classes=["large-output"]
                )
                gen_beginning_button = gr.Button("生成开头", visible=False)
                
            with gr.Tab("正文"):
                writing_plan_text = gr.Textbox(
                    "",
                    label="写作计划",
                    lines=5,
                    interactive=True
                )
                temp_setting_text = gr.Textbox(
                    "",
                    label="临时设定",
                    lines=5,
                    interactive=True
                )
                writing_memory_text = gr.Textbox(
                    "",
                    label="写作记忆",
                    lines=5,
                    interactive=True
                )
                novel_content_text = gr.Textbox(
                    "",
                    label="小说正文",
                    lines=15,
                    interactive=True,
                    elem_classes=["large-output"]
                )
                with gr.Row():
                    gen_next_paragraph_button = gr.Button("生成下一段", visible=False)
                    gen_ending_button = gr.Button("生成结局", visible=False)
                    download_novel_button = gr.Button("下载小说", visible=False)
        
        with gr.Column(scale=2, elem_id="row2"):
            gr.Markdown("### 创作对话记录")
            chat_history = gr.Chatbot(
                [],
                label="对话历史",
                elem_id="chatbot",
                height=600,
                show_label=False,
                elem_classes=["scrollable-chat"],
                bubble_full_width=True
            )
    
    # 公文创作界面
    with gr.Row(visible=False) as document_interface:
        with gr.Column(scale=1):
            with gr.Tab("公文生成"):
                doc_type = gr.Dropdown(
                    ["通知", "请示", "函", "报告"],
                    label="公文类型",
                    value="通知"
                )
                doc_requirements = gr.Textbox(
                    placeholder="请描述您需要生成的公文内容要求",
                    label="内容要求",
                    lines=10,
                    interactive=True,
                    elem_id="document_requirements"
                )
                stop_voice_button_4 = gr.Button(value="录音完毕")
                stop_voice_button_4.click(
                    stop_voice_recording,
                    inputs=[gr.Textbox(value="document_requirements", visible=False)],
                    outputs=[doc_requirements]
                )
                voice_input_button = gr.Button(value="语音输入")
                voice_input_button.click(
                    handle_voice_input,
                    inputs=[gr.Textbox(value="document_requirements", visible=False)],
                    outputs=[doc_requirements]
                )

                process_doc_button = gr.Button("生成公文", variant="primary")
                
            with gr.Tab("公文预览"):
                doc_preview = gr.Textbox(
                    "",
                    label="公文预览",
                    lines=20,
                    interactive=True,
                    elem_classes=["document-preview"]
                )
                download_doc_button = gr.Button("下载公文", visible=False)
        
        with gr.Column(scale=1):
            gr.Markdown("### 公文对话记录")
            doc_chat_history = gr.Chatbot(
                [],
                label="公文对话",
                height=600,
                show_label=False
            )
    
    # 剧本创作界面
    with gr.Row(visible=False) as script_interface:
        with gr.Column(scale=1):
            with gr.Tab("创意构思"):
                script_user_idea = gr.Textbox(
                    placeholder="请描述您的剧本创意",
                    label="创意构思",
                    lines=7,
                    interactive=True,
                    elem_id="script_idea"
                )
                stop_voice_button_5 = gr.Button(value="录音完毕")
                stop_voice_button_5.click(
                    stop_voice_recording,
                    inputs=[gr.Textbox(value="script_idea", visible=False)],
                    outputs=[script_user_idea]
                )
                voice_input_button = gr.Button(value="语音输入")
                voice_input_button.click(
                        handle_voice_input,
                        inputs=[gr.Textbox(value="script_idea", visible=False)],
                        outputs=[script_user_idea]
                    )

                with gr.Accordion("场景设置", open=False):
                    scene_count = gr.Number(
                        label="场景数量",
                        value=8,
                        minimum=3,
                        maximum=50,
                        step=1,
                        interactive=True
                    )
                    word_count_per_scene = gr.Number(
                        label="每场景字数",
                        value=800,
                        minimum=300,
                        maximum=5000,
                        step=100,
                        interactive=True
                    )
                script_user_requirements = gr.Textbox(
                    "",
                    label="剧本要求",
                    lines=7,
                    interactive=True,
                    elem_id="script_requirements" 
                )
                stop_voice_button_6 = gr.Button(value="录音完毕")
                stop_voice_button_6.click(
                    stop_voice_recording,
                    inputs=[gr.Textbox(value="script_requirements", visible=False)],
                    outputs=[script_user_requirements]
                )
                voice_input_button = gr.Button(value="语音输入")
                voice_input_button.click(
                        handle_voice_input,
                        inputs=[gr.Textbox(value="script_requirements", visible=False)],
                        outputs=[script_user_requirements]
                    )

                gen_script_outline_button = gr.Button("生成剧本大纲", variant="primary")
                
            with gr.Tab("剧本大纲"):
                script_outline_text = gr.Textbox(
                    "",
                    label="剧本大纲",
                    lines=15,
                    interactive=True
                )
                gen_script_beginning_button = gr.Button("生成", visible=False)
                
            with gr.Tab("剧本正文"):
                script_working_plan = gr.Textbox(
                    "",
                    label="工作计划",
                    lines=5,
                    interactive=True
                )
                script_temp_setting = gr.Textbox(
                    "",
                    label="临时设定",
                    lines=5,
                    interactive=True
                )
                script_memory = gr.Textbox(
                    "",
                    label="创作记忆",
                    lines=5,
                    interactive=True
                )
                script_content = gr.Textbox(
                    "",
                    label="剧本内容",
                    lines=15,
                    interactive=True,
                    elem_classes=["script-content"]
                )
                with gr.Row():
                    gen_next_scene_button = gr.Button("生成下一场", visible=False)
                    gen_script_ending_button = gr.Button("生成结局", visible=False)
                    download_script_button = gr.Button("下载剧本", visible=False)
        
        with gr.Column(scale=1):
            gr.Markdown("### 剧本创作对话")
            script_chat_history = gr.Chatbot(
                [],
                label="剧本对话",
                height=600,
                show_label=False,
                elem_classes=["script-panel"]
            )

    # 小说创作事件绑定
    gen_ouline_button.click(
        gen_ouline_button_clicked,
        [aign, user_idea_text, chat_history, chapter_count, word_count_per_chapter],
        [aign, chat_history, novel_outline_text, gen_beginning_button, gen_next_paragraph_button, gen_ending_button, download_novel_button]
    )
    
    gen_beginning_button.click(
        gen_beginning_button_clicked,
        [aign, chat_history, novel_outline_text, user_requriments_text, embellishment_idea_text],
        [aign, chat_history, writing_plan_text, temp_setting_text, novel_content_text, gen_next_paragraph_button, gen_ending_button, download_novel_button]
    )
    
    gen_next_paragraph_button.click(
        gen_next_paragraph_button_clicked,
        [aign, chat_history, user_idea_text, novel_outline_text, writing_memory_text, temp_setting_text, writing_plan_text, user_requriments_text, embellishment_idea_text],
        [aign, chat_history, writing_plan_text, temp_setting_text, writing_memory_text, novel_content_text, gen_next_paragraph_button, gen_ending_button, download_novel_button]
    )
    
    gen_ending_button.click(
        gen_ending_button_clicked,
        [aign, chat_history, user_requriments_text, novel_content_text],
        [aign, chat_history, writing_plan_text, temp_setting_text, writing_memory_text, novel_content_text, gen_next_paragraph_button, gen_ending_button, download_novel_button]
    )
    
    download_novel_button.click(
        download_novel,
        [],
        [download_novel_button]
    )
    
    # 公文创作事件绑定
    process_doc_button.click(
        process_document_request,
        [doc_type, doc_requirements, doc_chat_history],
        [doc_chat_history, doc_preview, download_doc_button]
    )
    
    download_doc_button.click(
        download_document,
        [doc_type, doc_preview],
        [download_doc_button]
    )
    
    # 剧本创作事件绑定
    gen_script_outline_button.click(
        gen_script_outline_button_clicked,
        [script_aign, script_user_idea, script_chat_history, scene_count, word_count_per_scene],
        [script_aign, script_chat_history, script_outline_text, gen_script_beginning_button]
    )
    
    gen_script_beginning_button.click(
        gen_script_beginning_button_clicked,
        [script_aign, script_chat_history, script_outline_text, script_user_requirements],
        [script_aign, script_chat_history, script_working_plan, script_temp_setting, script_content, gen_next_scene_button]
    )
    
    gen_next_scene_button.click(
        gen_next_scene_button_clicked,
        [script_aign, script_chat_history, script_user_idea, script_outline_text, script_memory, script_temp_setting, script_working_plan, script_user_requirements],
        [script_aign, script_chat_history, script_working_plan, script_temp_setting, script_memory, script_content]
    )
    
    gen_script_ending_button.click(
        gen_script_ending_button_clicked,
        [script_aign, script_chat_history, script_user_requirements, script_content],
        [script_aign, script_chat_history, script_working_plan, script_temp_setting, script_memory, script_content, download_script_button]
    )
    
    download_script_button.click(
        download_script,
        [],
        [download_script_button]
    )
    
    # 界面切换事件
    switch_btn.click(
        lambda x: "document" if x == "fiction" else ("script" if x == "document" else "fiction"),
        [current_mode],
        [current_mode],
    ).then(
        toggle_interface,
        [current_mode],
        [fiction_interface, document_interface, script_interface, switch_btn]
    )
    
    # 修改按钮绑定代码
    voice_input_button.click(
        handle_voice_input,
        [user_idea_text],
        [user_idea_text]
    )

    stop_voice_button_1.click(
        stop_voice_recording,
        [user_idea_text],
        [user_idea_text]
    )

    stop_voice_button_2.click(
        stop_voice_recording,
        [user_requriments_text],
        [user_requriments_text]
    )

    stop_voice_button_3.click(
        stop_voice_recording,
        [embellishment_idea_text],
        [embellishment_idea_text]
    )

    stop_voice_button_4.click(
        stop_voice_recording,
        [doc_requirements],
        [doc_requirements]
    )

    stop_voice_button_5.click(
        stop_voice_recording,
        [script_user_idea],
        [script_user_idea]
    )

    stop_voice_button_6.click(
        stop_voice_recording,
        [script_user_requirements],
        [script_user_requirements]
    )

if __name__ == "__main__":
    demo.launch(
        server_name=os.environ.get("SERVER_NAME", "127.0.0.1"),
        server_port=int(os.environ.get("SERVER_PORT", "7860")),
        share=os.environ.get("GRADIO_SHARE", "false").lower() == "true",
    )
